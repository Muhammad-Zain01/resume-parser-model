#!/usr/bin/env python3
"""Extract PDF/DOCX text and mark image-based resumes for later OCR."""

from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
import re
from dataclasses import dataclass
from pathlib import Path
import unicodedata
import zipfile

from docx import Document
from pypdf import PdfReader
from tqdm import tqdm


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DATA_DIR = PROJECT_ROOT / ".data"
IMAGE_BASED_MARKER = "[IMAGE_BASED]"
NO_TEXT_MARKER = "[NO_TEXT]"
EXTRACTION_FAILED_MARKER = "[EXTRACTION_FAILED]"


@dataclass(frozen=True)
class ExtractionRecord:
    """Result metadata for one processed resume."""

    resume_id: str
    category: str
    file_type: str
    raw_path: str
    extracted_path: str
    extraction_status: str
    page_count: int
    character_count: int
    extraction_error: str = ""


@dataclass(frozen=True)
class ExtractionResult:
    """Normalized result returned by the single document extractor."""

    text: str
    file_type: str
    status: str
    page_count: int = 0
    image_count: int = 0


def clean_extracted_text(text: str) -> str:
    """Apply safe cleanup while preserving resume structure and punctuation."""
    normalized = unicodedata.normalize("NFKC", text)
    # Some PDF fonts produce lone UTF-16 surrogate characters. Replace those
    # before writing UTF-8 text files so one malformed glyph cannot fail a job.
    normalized = normalized.encode("utf-8", errors="replace").decode("utf-8")
    normalized = normalized.replace("\u00a0", " ")
    normalized = re.sub(r"[\u200b-\u200d\ufeff]", "", normalized)
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")

    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in normalized.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _extract_pdf(path: Path) -> ExtractionResult:
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = clean_extracted_text("\n\n".join(page for page in pages if page))
    status = "TEXT_EXTRACTED" if text else "IMAGE_BASED"
    return ExtractionResult(
        text=text,
        file_type="PDF",
        status=status,
        page_count=len(reader.pages),
    )


def _extract_docx(path: Path) -> ExtractionResult:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)

    text = clean_extracted_text("\n".join(parts))
    image_count = len(document.inline_shapes)
    if text:
        status = "TEXT_EXTRACTED"
    elif image_count:
        status = "IMAGE_BASED"
    else:
        status = "NO_TEXT"

    return ExtractionResult(
        text=text,
        file_type="DOCX",
        status=status,
        image_count=image_count,
    )


def _detect_file_type(path: Path) -> str:
    """Detect PDF/DOCX by file signature, including mislabeled DOCX files."""
    with path.open("rb") as handle:
        signature = handle.read(8)

    if signature.startswith(b"%PDF"):
        return "PDF"

    if signature.startswith(b"PK"):
        try:
            with zipfile.ZipFile(path) as archive:
                if "word/document.xml" in archive.namelist():
                    return "DOCX"
        except zipfile.BadZipFile:
            pass

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "PDF"
    if suffix == ".docx":
        return "DOCX"
    raise ValueError(f"Unsupported document type: {path.suffix}")


def extract_document(path: Path) -> ExtractionResult:
    """Extract and safely clean one PDF or DOCX document."""
    file_type = _detect_file_type(path)
    if file_type == "PDF":
        return _extract_pdf(path)
    if file_type == "DOCX":
        return _extract_docx(path)
    raise ValueError(f"Unsupported document type: {path.suffix}")


class ResumeTextExtractor:
    """Extract text while preserving the raw category directory structure."""

    def __init__(
        self,
        raw_dir: Path | None = None,
        extracted_dir: Path | None = None,
        overwrite: bool = False,
        workers: int = 8,
    ) -> None:
        self.raw_dir = raw_dir or self._default_raw_dir()
        self.extracted_dir = extracted_dir or DATA_DIR / "extracted"
        self.overwrite = overwrite
        if workers < 1:
            raise ValueError("workers must be at least 1")
        self.workers = workers

    @staticmethod
    def _default_raw_dir() -> Path:
        """Use ``raw`` when available and support the current ``pdf`` name."""
        raw_dir = DATA_DIR / "raw"
        return raw_dir if raw_dir.is_dir() else DATA_DIR / "pdf"

    def find_documents(self) -> list[Path]:
        if not self.raw_dir.is_dir():
            raise FileNotFoundError(f"Raw resume directory does not exist: {self.raw_dir}")
        return sorted(
            path
            for path in self.raw_dir.rglob("*")
            if path.is_file() and path.suffix.lower() in {".pdf", ".docx"}
        )

    def _category_for(self, document_path: Path) -> str:
        parent = document_path.relative_to(self.raw_dir).parent
        return str(parent) if str(parent) != "." else "Uncategorized"

    def _text_path_for(self, document_path: Path) -> Path:
        relative_path = document_path.relative_to(self.raw_dir).with_suffix(".txt")
        return self.extracted_dir / relative_path

    @staticmethod
    def _relative_path(path: Path) -> str:
        return path.resolve().relative_to(PROJECT_ROOT).as_posix()

    @staticmethod
    def _write_text(path: Path, content: str) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")

    def _process_document(self, document_path: Path) -> ExtractionRecord:
        text_path = self._text_path_for(document_path)
        category = self._category_for(document_path)

        try:
            result = extract_document(document_path)

            # Existing files are preserved so a future OCR result is not lost.
            if not text_path.exists() or self.overwrite:
                if result.text:
                    content = result.text
                elif result.status == "IMAGE_BASED":
                    content = IMAGE_BASED_MARKER
                else:
                    content = NO_TEXT_MARKER
                self._write_text(text_path, content)

            current_text = text_path.read_text(encoding="utf-8", errors="replace")
            return ExtractionRecord(
                resume_id=document_path.stem,
                category=category,
                file_type=result.file_type,
                raw_path=self._relative_path(document_path),
                extracted_path=self._relative_path(text_path),
                extraction_status=result.status,
                page_count=result.page_count,
                character_count=len(current_text),
            )
        except Exception as error:  # Continue processing if one document is damaged.
            self._write_text(text_path, EXTRACTION_FAILED_MARKER)
            return ExtractionRecord(
                resume_id=document_path.stem,
                category=category,
                file_type=document_path.suffix.lower().lstrip(".").upper(),
                raw_path=self._relative_path(document_path),
                extracted_path=self._relative_path(text_path),
                extraction_status="EXTRACTION_FAILED",
                page_count=0,
                character_count=len(EXTRACTION_FAILED_MARKER),
                extraction_error=str(error),
            )

    def run(self) -> list[ExtractionRecord]:
        """Process all supported documents and return extraction results."""
        document_paths = self.find_documents()
        with ThreadPoolExecutor(max_workers=self.workers) as executor:
            results = executor.map(self._process_document, document_paths)
            return list(
                tqdm(
                    results,
                    total=len(document_paths),
                    desc=f"Extracting resumes ({self.workers} workers)",
                    unit="document",
                )
            )

    def find_image_based_texts(self) -> list[Path]:
        """Return text files that still contain the image-based marker."""
        if not self.extracted_dir.is_dir():
            return []
        return sorted(
            path
            for path in self.extracted_dir.rglob("*.txt")
            if IMAGE_BASED_MARKER in path.read_text(encoding="utf-8", errors="replace")
        )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path)
    parser.add_argument("--extracted-dir", type=Path)
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Replace existing text files; do not use after OCR without a backup.",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=8,
        help="Number of parallel document workers (default: 8).",
    )
    parser.add_argument(
        "--list-image-based",
        action="store_true",
        help="List text files containing [IMAGE_BASED] without extracting PDFs.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    extractor = ResumeTextExtractor(
        raw_dir=args.raw_dir,
        extracted_dir=args.extracted_dir,
        overwrite=args.overwrite,
        workers=args.workers,
    )

    if args.list_image_based:
        image_files = extractor.find_image_based_texts()
        for path in image_files:
            print(path)
        print(f"Image-based text files: {len(image_files)}")
        return 0

    records = extractor.run()
    counts: dict[str, int] = {}
    for record in records:
        counts[record.extraction_status] = counts.get(record.extraction_status, 0) + 1
    print(f"Processed documents: {len(records)}")
    file_types = {record.file_type for record in records}
    for file_type in sorted(file_types):
        print(f"{file_type}: {sum(record.file_type == file_type for record in records)}")
    for status, count in sorted(counts.items()):
        print(f"{status}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
